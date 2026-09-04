import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_proposal():
    doc = Document()
    
    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Color Palette
    HEX_PRIMARY = "1B365D"    # Deep Navy Blue
    HEX_SECONDARY = "008080"  # Ocean Teal
    HEX_DARK = "222222"       # Charcoal Body Text
    HEX_LIGHT_BG = "F0F4F8"   # Table alternating row / callout bg
    HEX_BORDER = "CCCCCC"     # Muted gray border
    
    COLOR_PRIMARY = RGBColor(0x1B, 0x36, 0x5D)
    COLOR_SECONDARY = RGBColor(0x00, 0x80, 0x80)
    COLOR_DARK = RGBColor(0x22, 0x22, 0x22)
    COLOR_MUTED = RGBColor(0x66, 0x66, 0x66)

    # Helper Functions
    def set_cell_background(cell, hex_color):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin, value in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{margin}')
            node.set(qn('w:w'), str(value))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_table_borders(table):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_PRIMARY}"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_custom_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        if level == 1:
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run.font.size = Pt(16)
            run.font.color.rgb = COLOR_PRIMARY
            run.font.name = "Calibri"
            # Add bottom accent line under H1
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="{HEX_PRIMARY}"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
        elif level == 2:
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run.font.size = Pt(13)
            run.font.color.rgb = COLOR_SECONDARY
            run.font.name = "Calibri"
        elif level == 3:
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run.font.size = Pt(11)
            run.font.color.rgb = COLOR_DARK
            run.font.name = "Calibri"
        return p

    def add_body_paragraph(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.size = Pt(10.5)
            r_bold.font.color.rgb = COLOR_DARK
            r_bold.font.name = "Calibri"
            
        r_text = p.add_run(text)
        r_text.italic = italic
        r_text.font.size = Pt(10.5)
        r_text.font.color.rgb = COLOR_DARK
        r_text.font.name = "Calibri"
        return p

    def add_bullet_item(bold_title, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        r_bold = p.add_run(bold_title + " ")
        r_bold.bold = True
        r_bold.font.size = Pt(10.5)
        r_bold.font.color.rgb = COLOR_DARK
        r_bold.font.name = "Calibri"
        
        r_text = p.add_run(text)
        r_text.font.size = Pt(10.5)
        r_text.font.color.rgb = COLOR_DARK
        r_text.font.name = "Calibri"
        return p

    def add_callout_box(text, title="NOTE:"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, HEX_LIGHT_BG)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        # Thick left border
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_PRIMARY}"/>'
            f'<w:top w:val="none"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r_title = p.add_run(f"{title} ")
        r_title.bold = True
        r_title.font.size = Pt(10)
        r_title.font.color.rgb = COLOR_PRIMARY
        r_title.font.name = "Calibri"
        
        r_text = p.add_run(text)
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = COLOR_DARK
        r_text.font.name = "Calibri"
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # COVER / HEADER BLOCK
    # -------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("COMMERCIAL & TECHNICAL PROPOSAL")
    run_title.bold = True
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = COLOR_PRIMARY
    run_title.font.name = "Calibri"

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("AutoBell / IoBell Smart Audio & Automated Bell Management System")
    run_sub.bold = True
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = COLOR_SECONDARY
    run_sub.font.name = "Calibri"

    # Meta Table (Prepared For / Prepared By)
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table)
    
    meta_data = [
        [("PREPARED FOR:", "Valued Client / Educational Institution"), ("PROPOSAL DATE:", "August 2026")],
        [("PREPARED BY:", "Digitap Business Solution (digitapbs.pk)"), ("DOCUMENT REF:", "DBS-AUTOBELL-PROP-2026")]
    ]
    
    for row_idx, row in enumerate(meta_table.rows):
        for col_idx, cell in enumerate(row.cells):
            set_cell_background(cell, HEX_LIGHT_BG)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            lbl, val = meta_data[row_idx][col_idx]
            r_lbl = p.add_run(lbl + " ")
            r_lbl.bold = True
            r_lbl.font.size = Pt(9.5)
            r_lbl.font.color.rgb = COLOR_PRIMARY
            r_val = p.add_run(val)
            r_val.font.size = Pt(9.5)
            r_val.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # SECTION 1: EXECUTIVE SUMMARY
    # -------------------------------------------------------------
    add_custom_heading("1. Executive Summary", level=1)
    add_body_paragraph(
        "AutoBell (IoBell) is an enterprise-grade, cloud-managed IoT smart bell and campus audio broadcast platform developed by Digitap Business Solution. Designed specifically for educational institutions (Schools, Colleges, Universities, and Madrasas), AutoBell eliminates manual bell errors, hardware freezes, and complex timer box programming by providing seamless cloud control via Web Dashboard and Mobile App."
    )
    add_body_paragraph(
        "This proposal outlines a turnkey solution combining hardware integration, custom software modifications, audio expansion via extra speakers, high-performance power amplification, and professional on-site wiring and installation."
    )

    add_callout_box(
        "AutoBell features a built-in Hardware Watchdog and local SPIFFS/LittleFS caching. Even during complete internet or network outages, all scheduled bells continue to ring with 100% precision.",
        title="KEY ADVANTAGE:"
    )

    # -------------------------------------------------------------
    # SECTION 2: SYSTEM ARCHITECTURE & CORE CAPABILITIES
    # -------------------------------------------------------------
    add_custom_heading("2. Core System Architecture", level=1)
    add_body_paragraph("The AutoBell platform consists of three seamlessly integrated tiers:")
    
    add_bullet_item("1. ESP32 Smart IoT Controller:", "On-site industrial-grade hardware controller connected to PA amplifiers and bell relays, equipped with dual-core processing, Wi-Fi/Ethernet sync, and local offline cache storage.")
    add_bullet_item("2. Cloud Web Dashboard & Mobile App:", "Central management platform accessible from any browser or smartphone (Android/iOS) for single-tap schedule switching, live audio broadcasting, and emergency alerts.")
    add_bullet_item("3. Automated Audio & TTS Engine:", "Supports pre-announcement audio jingles (with configurable 2–5 second delays), natural Text-To-Speech (TTS) voice notices, and custom MP3 file execution.")

    # -------------------------------------------------------------
    # SECTION 3: CUSTOM SOFTWARE CHANGES & TWEAKS
    # -------------------------------------------------------------
    add_custom_heading("3. Custom Software Changes & Tweaks", level=1)
    add_body_paragraph(
        "To ensure AutoBell perfectly integrates into your institution's specific daily routines and administrative workflows, this proposal includes customized software developments and dashboard tweaks:"
    )
    
    add_bullet_item("Tailored Bell Schedule Profiles:", "Pre-configuration of multi-tier bell profiles including Normal Day, Exam Schedule, Ramadan Timing, Friday Schedule, and Half-Day events with 1-click dashboard switching.")
    add_bullet_item("Custom Pre-Announcement Jingle & Delay Logic:", "Integration of custom institutional audio chimes and tailored delay timing (custom 2s, 3s, or 5s pauses) prior to bell ringing or voice broadcasts.")
    add_bullet_item("Institution Branding & Custom Subdomain:", "Customization of the Web Dashboard interface featuring your school logo, institutional color theme, and dedicated sub-domain access.")
    add_bullet_item("Role-Based Access Control (RBAC):", "Custom security hierarchy setup for Principal, Vice Principal, IT Administrator, and Staff, ensuring restricted authorization for emergency lockdown overrides.")
    add_bullet_item("Custom Text-to-Speech (TTS) Voice Tuning:", "Optimization of automated speech parameters (accent, pitch, speed, and multi-lingual voice engine) tailored for clear audibility over campus speakers.")
    add_bullet_item("API / Webhook SMS Integration:", "Optional backend webhook configuration to automatically trigger instant SMS alerts to management when an emergency lockdown bell is activated.")

    # -------------------------------------------------------------
    # SECTION 4: EXTRA SPEAKERS & AUDIO EXPANSION
    # -------------------------------------------------------------
    add_custom_heading("4. Extra Speakers & Audio Expansion", level=1)
    add_body_paragraph(
        "To guarantee crystal-clear audibility across all campus zones—including open grounds, corridors, assembly halls, and administrative wings—we offer high-output extra speaker deployment:"
    )

    add_bullet_item("Heavy-Duty Weatherproof Horn Speakers (Outdoor):", "IP66 weather-sealed aluminum/ABS horn speakers (40W–50W) designed for wide sound dispersion in outdoor playfields, main gates, and open assembly areas.")
    add_bullet_item("Indoor Wall & Column Speakers (Hallways & Classrooms):", "Aesthetically styled 20W–30W indoor column/box speakers engineered for high vocal clarity without harsh resonance in corridors and academic blocks.")
    add_bullet_item("Multi-Zone PA Audio Power Amplifier:", "High-performance commercial PA amplifier (250W–500W output) equipped with multi-zone speaker selection, line inputs for AutoBell ESP32 controller, aux inputs, and thermal overload protection.")
    add_bullet_item("Acoustic Balancing & Impedance Matching:", "Custom matching transformer setup (70V/100V line system) to ensure uniform audio levels without signal degradation over long cable distances.")

    # -------------------------------------------------------------
    # SECTION 5: WIRING, MATERIALS & INSTALLATION CHARGES
    # -------------------------------------------------------------
    add_custom_heading("5. Wiring, Materials & Installation Charges", level=1)
    add_body_paragraph(
        "Proper wiring infrastructure and professional installation are essential for audio clarity, equipment longevity, and campus safety. Our complete turnkey installation package covers:"
    )

    add_bullet_item("Pure Copper Heavy-Duty Audio Cabling:", "2-core flexible copper audio wire designed for minimal resistance and high-fidelity sound transmission over extended campus lengths.")
    add_bullet_item("Protective PVC Conduits & Accessories:", "High-impact PVC piping, trunking ducts, flexible conduits, saddle clamps, junction boxes, and connectors to ensure safe and hidden wiring.")
    add_bullet_item("Mounting Hardware & Fixtures:", "Heavy-duty steel brackets, expansion bolts, wall plugs, and vibration-dampening mounts for all indoor and outdoor speakers.")
    add_bullet_item("On-Site Professional Installation Services:", 
                    "\n• Complete cable laying and conduit routing across campus blocks."
                    "\n• Solid mounting and directional alignment of all extra horn and column speakers."
                    "\n• Rack/Wall mounting of AutoBell Controller unit and Power Amplifier."
                    "\n• Integration with existing PA system or new amplifier units."
                    "\n• Wi-Fi / LAN setup, cloud pairing, and offline SPIFFS schedule synchronization."
                    "\n• Multi-zone sound testing, acoustic balancing, and volume calibration."
                    "\n• Complete hands-on training for school staff and IT administration.")

    # -------------------------------------------------------------
    # SECTION 6: FINANCIAL PROPOSAL & COST BREAKDOWN
    # -------------------------------------------------------------
    add_custom_heading("6. Financial Proposal & Cost Breakdown", level=1)
    add_body_paragraph(
        "The following table details the turnkey investment for AutoBell hardware, custom software development, extra speakers, wiring materials, and complete installation charges:"
    )

    fin_table = doc.add_table(rows=1, cols=5)
    fin_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(fin_table)

    headers = ["Item / Description", "Category", "Qty", "Unit Price (PKR)", "Total (PKR)"]
    hdr_cells = fin_table.rows[0].cells
    for i, title in enumerate(headers):
        set_cell_background(hdr_cells[i], HEX_PRIMARY)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i in [2, 3, 4] else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.name = "Calibri"

    items = [
        ("AutoBell Smart IoT Controller Unit (ESP32 + Dual Relays + Line Out)", "Hardware", "1 Unit", "35,000", "35,000"),
        ("AutoBell Cloud SaaS License & Mobile App Access (1-Year Subscription)", "Software", "1 Year", "25,000", "25,000"),
        ("Custom Software Changes & Tweaks (Profiles, TTS Tuning, Custom Jingle & Branding)", "Software Customization", "1 Service", "15,000", "15,000"),
        ("Heavy-Duty Weatherproof Outdoor Horn Speakers (40W-50W)", "Speakers", "4 Units", "8,500", "34,000"),
        ("Indoor Wall / Column Speakers (20W-30W for Corridors)", "Speakers", "4 Units", "5,500", "22,000"),
        ("Commercial PA Power Amplifier (350W Multi-Zone with Line In)", "Amplification", "1 Unit", "32,000", "32,000"),
        ("Pure Copper Audio Wiring & Heavy PVC Conduit Piping Pack", "Wiring & Materials", "1 Job", "18,000", "18,000"),
        ("On-Site Professional Installation, Speaker Mounting, Calibration & Staff Training", "Installation Labor", "1 Job", "20,000", "20,000"),
    ]

    total_amount = 201000

    for idx, row_data in enumerate(items):
        row = fin_table.add_row()
        bg_color = HEX_LIGHT_BG if idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if c_idx in [2, 3, 4] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            r.font.color.rgb = COLOR_DARK
            r.font.name = "Calibri"

    # Total Row
    tot_row = fin_table.add_row()
    for c_idx in range(5):
        cell = tot_row.cells[c_idx]
        set_cell_background(cell, HEX_PRIMARY)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        if c_idx == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run("GRAND TOTAL INVESTMENT (TAXES INCLUDED)")
            r.bold = True
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        elif c_idx == 4:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(f"PKR {total_amount:,}")
            r.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_body_paragraph(
        "Commercial Terms & Payment Conditions:\n"
        "• Payment Terms: 50% advance upon purchase order sign-off, 40% upon hardware delivery and wiring completion, 10% post-installation testing and sign-off.\n"
        "• Proposal Validity: Valid for 30 days from document issuance date.\n"
        "• Delivery & Installation Lead Time: 5 to 7 working days from advance payment confirmation."
    )

    # -------------------------------------------------------------
    # SECTION 7: WARRANTY, SLA & MAINTENANCE
    # -------------------------------------------------------------
    add_custom_heading("7. Warranty, SLA & Support Framework", level=1)
    add_bullet_item("1-Year Hardware Replacement SLA:", "Comprehensive 1-year replacement warranty on the ESP32 AutoBell IoT Controller against manufacturing defects.")
    add_bullet_item("Continuous Cloud Health Monitoring:", "Automated cloud ping diagnostics and hardware watchdog alerts to detect hardware or network disconnections proactively.")
    add_bullet_item("Free Over-The-Air (OTA) Firmware Updates:", "Automated cloud firmware upgrades providing new features, security patches, and system improvements at zero extra cost.")
    add_bullet_item("Dedicated Technical Support:", "Direct phone, email, and WhatsApp ticketing support (Monday – Saturday, 8:00 AM – 6:00 PM).")

    # -------------------------------------------------------------
    # SECTION 8: ACCEPTANCE & SIGN-OFF BLOCK
    # -------------------------------------------------------------
    add_custom_heading("8. Proposal Acceptance & Sign-off", level=1)
    add_body_paragraph(
        "By signing below, the Client approves the technical scope, custom software tweaks, extra speakers, installation charges, and commercial terms outlined in this proposal."
    )

    sig_table = doc.add_table(rows=4, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(sig_table)

    sig_data = [
        [("FOR CLIENT:", ""), ("FOR DIGITAP BUSINESS SOLUTION:", "")],
        [("Authorized Signature: __________________", ""), ("Authorized Signature: __________________", "")],
        [("Name: _______________________________", ""), ("Name: _______________________________", "")],
        [("Designation: ________________________", ""), ("Designation: ________________________", "")],
    ]

    for r_idx, row in enumerate(sig_table.rows):
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, HEX_LIGHT_BG if r_idx == 0 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            txt1, txt2 = sig_data[r_idx][c_idx]
            r1 = p.add_run(txt1)
            r1.bold = (r_idx == 0)
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = COLOR_PRIMARY if r_idx == 0 else COLOR_DARK

    # Save document
    output_path = r"c:\My Drive\My Drive\AutoBell\AutoBell_Commercial_Proposal.docx"
    doc.save(output_path)
    print(f"Proposal Word document successfully created at: {output_path}")

if __name__ == "__main__":
    create_proposal()

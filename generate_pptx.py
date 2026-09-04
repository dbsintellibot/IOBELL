import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Constants
    COLOR_NAVY = RGBColor(15, 23, 42)      # #0F172A - Primary Dark
    COLOR_SLATE_BG = RGBColor(248, 250, 252) # #F8FAFC - Main Slide BG
    COLOR_BLUE_ACCENT = RGBColor(37, 99, 235) # #2563EB - Royal Blue Accent
    COLOR_CYAN = RGBColor(2, 132, 199)     # #0284C7 - Cyan Accent
    COLOR_GOLD = RGBColor(245, 158, 11)    # #F59E0B - Gold/Amber Highlight
    COLOR_WHITE = RGBColor(255, 255, 255)  # White
    COLOR_DARK_TEXT = RGBColor(30, 41, 59) # #1E293B
    COLOR_MUTED_TEXT = RGBColor(100, 116, 139) # #64748B
    COLOR_CARD_BG = RGBColor(255, 255, 255)
    COLOR_CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0
    COLOR_CARD_HEADER_BG = RGBColor(241, 245, 249) # #F1F5F9
    COLOR_RED_ACCENT = RGBColor(225, 29, 72) # #E11D48 - Emergency Red

    # Helper: Set Solid Background Color
    def set_bg(slide, color):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color

    # Helper: Add Standard Header to Content Slides
    def add_header(slide, category, title, dark_mode=False):
        # Category Tracker / Eyebrow
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.3))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_GOLD if dark_mode else COLOR_BLUE_ACCENT
        p_cat.font.name = "Arial"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.5), Inches(0.6))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_WHITE if dark_mode else COLOR_NAVY
        p_title.font.name = "Arial"

        # Decorative Accent Bar below header
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(11.733), Inches(0.04))
        bar.fill.solid()
        bar.fill.fore_color.rgb = COLOR_GOLD if dark_mode else COLOR_BLUE_ACCENT
        bar.line.fill.background()

    # Helper: Add Footer
    def add_footer(slide, current_slide, total_slides=12, dark_mode=False):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.3))
        tf = footer_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"AutoBell | Next-Gen Smart Campus Bell & Audio System   •   https://iobell.web.app/   •   Slide {current_slide} of {total_slides}"
        p.font.size = Pt(9)
        p.font.color.rgb = COLOR_MUTED_TEXT if not dark_mode else RGBColor(148, 163, 184)
        p.font.name = "Arial"

    # Helper: Add Card Shape
    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    # ==========================================
    # SLIDE 1: TITLE SLIDE (Dark Theme)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_bg(slide1, COLOR_NAVY)

    # Decorative Top Badge
    badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(3.2), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(30, 41, 59)
    badge.line.color.rgb = COLOR_BLUE_ACCENT
    badge.line.width = Pt(1.5)
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "⚡ NEXT-GEN IOT SAAS PLATFORM"
    p_b.alignment = PP_ALIGN.CENTER
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_GOLD

    # Main Title Box
    tbox = slide1.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(7.5), Inches(2.2))
    tf = tbox.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "AutoBell"
    p1.font.size = Pt(54)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.font.name = "Arial"

    p2 = tf.add_paragraph()
    p2.text = "Smart Campus Bell & Audio Management System"
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_CYAN
    p2.font.name = "Arial"

    # Subtitle / Tagline
    sbox = slide1.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(7.5), Inches(1.2))
    stf = sbox.text_frame
    stf.word_wrap = True
    sp1 = stf.paragraphs[0]
    sp1.text = "Modernizing school operations, automated announcements, and campus emergency broadcasting — with 100% cloud control and offline-first reliability."
    sp1.font.size = Pt(15)
    sp1.font.color.rgb = RGBColor(203, 213, 225)

    # Add Product Image on Right side if available
    img_path = r"c:\My Drive\My Drive\AutoBell\new_autobell_picture-removebg-preview.png"
    if os.path.exists(img_path):
        slide1.shapes.add_picture(img_path, Inches(8.5), Inches(1.4), width=Inches(4.2))

    # Presenter / Info Box at bottom left
    infobox = slide1.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(7.5), Inches(0.8))
    itf = infobox.text_frame
    ip = itf.paragraphs[0]
    ip.text = "🌐 Live Dashboard: iobell.web.app   |   🏢 Enterprise Customer Presentation"
    ip.font.size = Pt(12)
    ip.font.bold = True
    ip.font.color.rgb = COLOR_GOLD

    add_footer(slide1, 1, 12, dark_mode=True)

    # ==========================================
    # SLIDE 2: EXECUTIVE SUMMARY & PROBLEM STATEMENT
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_bg(slide2, COLOR_SLATE_BG)
    add_header(slide2, "Executive Overview", "The Modern Campus Challenge vs. The AutoBell Solution")

    # Left Card: Traditional Problems
    add_card(slide2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1), RGBColor(254, 242, 242), RGBColor(254, 202, 202))
    
    t_left = slide2.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.7))
    tf_l = t_left.text_frame
    tf_l.word_wrap = True
    
    pl = tf_l.paragraphs[0]
    pl.text = "❌ Outdated Legacy Bell Systems"
    pl.font.size = Pt(18)
    pl.font.bold = True
    pl.font.color.rgb = COLOR_RED_ACCENT
    
    items_l = [
        "Manual & Error-Prone Timer Boxes: Complex button programming required on-site.",
        "Rigid Schedule Changes: Switching to Exam Mode, Ramadan, or Half-Days causes chaos.",
        "System Freezes & Outages: Timer boxes freeze unexpectedly, skipping critical bells.",
        "No Voice / Speech Capabilities: Bell timers can only ring simple gongs, no spoken notices.",
        "Slow Emergency Response: Manual staff action required during campus lockdown emergencies."
    ]
    for item in items_l:
        p = tf_l.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(8)

    # Right Card: AutoBell Solution
    add_card(slide2, Inches(6.933), Inches(1.6), Inches(5.6), Inches(5.1), RGBColor(240, 253, 244), RGBColor(187, 247, 208))
    
    t_right = slide2.shapes.add_textbox(Inches(7.133), Inches(1.8), Inches(5.2), Inches(4.7))
    tf_r = t_right.text_frame
    tf_r.word_wrap = True
    
    pr = tf_r.paragraphs[0]
    pr.text = "✨ The AutoBell Smart Solution"
    pr.font.size = Pt(18)
    pr.font.bold = True
    pr.font.color.rgb = RGBColor(22, 101, 52)
    
    items_r = [
        "100% Cloud Remote Control: Manage profiles and schedules from any smartphone or browser.",
        "1-Tap Profile Switching: Toggle Normal Day, Exam Mode, or Ramadan schedules instantly.",
        "Offline-First ESP32 Hardware: Local memory caching guarantees bells ring even without Wi-Fi.",
        "Text-to-Speech & Voice Broadcasts: Automated spoken announcements and live voice streaming.",
        "Instant Emergency Lockdown: Slide-to-activate lockdown siren & priority audio override."
    ]
    for item in items_r:
        p = tf_r.add_paragraph()
        p.text = "✔ " + item
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(8)

    add_footer(slide2, 2, 12)

    # ==========================================
    # SLIDE 3: CORE CAPABILITIES GRID (6 Cards)
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_bg(slide3, COLOR_SLATE_BG)
    add_header(slide3, "Product Features", "6 Core Pillars of the AutoBell Ecosystem")

    pillars = [
        ("📅 Cloud Profile Scheduling", "Create, edit, and instantly switch profiles (Normal Day, Exam, Ramadan, Weekend) remotely via Web or Mobile."),
        ("🎵 Pre-Announcement Jingle", "Play a signature campus chime followed by a configurable 2–5s pause before any bell, speech, or audio notice."),
        ("📢 TTS & Live Broadcasts", "Type any text for instant natural voice broadcast, or stream live/recorded MP3 audio directly to classrooms."),
        ("⚡ Offline-First Execution", "Schedules cache locally inside ESP32 SPIFFS memory. Bells ring on time 100% of the time, even offline."),
        ("🛡️ Hardware Watchdog Uptime", "Built-in automatic freeze recovery monitors device state continuously to ensure zero manual reboots."),
        ("🚨 Emergency Lockdown Mode", "Dedicated one-slide trigger on Web & Mobile apps for immediate, campus-wide emergency siren override.")
    ]

    card_coords = [
        (Inches(0.8), Inches(1.6)), (Inches(4.85), Inches(1.6)), (Inches(8.9), Inches(1.6)),
        (Inches(0.8), Inches(4.3)), (Inches(4.85), Inches(4.3)), (Inches(8.9), Inches(4.3))
    ]

    for idx, (title, desc) in enumerate(pillars):
        x, y = card_coords[idx]
        add_card(slide3, x, y, Inches(3.633), Inches(2.4), COLOR_WHITE, COLOR_CARD_BORDER)
        
        # Header Accent Stripe
        stripe = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(3.633), Inches(0.08))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = COLOR_BLUE_ACCENT if idx != 5 else COLOR_RED_ACCENT
        stripe.line.fill.background()

        tbox = slide3.shapes.add_textbox(x + Inches(0.15), y + Inches(0.2), Inches(3.333), Inches(2.0))
        tf = tbox.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY if idx != 5 else COLOR_RED_ACCENT
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_DARK_TEXT
        p2.space_before = Pt(8)

    add_footer(slide3, 3, 12)

    # ==========================================
    # SLIDE 4: CLOUD REMOTE SCHEDULING & PROFILES
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_bg(slide4, COLOR_SLATE_BG)
    add_header(slide4, "Schedule Automation", "Flexible Profile Management & Remote Cloud Sync")

    # Left Section: Profiles Overview
    add_card(slide4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1), COLOR_WHITE, COLOR_CARD_BORDER)
    
    t_p = slide4.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.7))
    tf_p = t_p.text_frame
    tf_p.word_wrap = True
    
    p = tf_p.paragraphs[0]
    p.text = "🎯 Multi-Profile Campus Scheduling"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_ACCENT

    prof_list = [
        "Normal Day Schedule: Standard lesson slots, recess, assembly, and dismissal bells.",
        "Exam Mode: Custom quiet timings, extended periods, and distinct warning chimes.",
        "Ramadan & Winter Timings: Adjusted shortened durations applied in 1 tap.",
        "Half-Day & Special Events: Instant override without overwriting primary master schedules.",
        "Weekend / Holiday Silence: Automatic suppression of all non-essential alarms."
    ]
    for item in prof_list:
        pi = tf_p.add_paragraph()
        pi.text = "▪ " + item
        pi.font.size = Pt(13)
        pi.font.color.rgb = COLOR_DARK_TEXT
        pi.space_before = Pt(10)

    # Right Section: Key Highlights
    add_card(slide4, Inches(6.933), Inches(1.6), Inches(5.6), Inches(2.4), RGBColor(239, 246, 255), RGBColor(191, 219, 254))
    tb_r1 = slide4.shapes.add_textbox(Inches(7.133), Inches(1.8), Inches(5.2), Inches(2.0))
    tf_r1 = tb_r1.text_frame
    tf_r1.word_wrap = True
    pr1 = tf_r1.paragraphs[0]
    pr1.text = "⚡ Instant Cloud Synchronization"
    pr1.font.size = Pt(16)
    pr1.font.bold = True
    pr1.font.color.rgb = COLOR_BLUE_ACCENT
    
    pr2 = tf_r1.add_paragraph()
    pr2.text = "Any schedule modified on the Web Dashboard or Mobile App syncs to the on-site ESP32 hardware in milliseconds via Supabase real-time sockets."
    pr2.font.size = Pt(13)
    pr2.font.color.rgb = COLOR_DARK_TEXT
    pr2.space_before = Pt(6)

    add_card(slide4, Inches(6.933), Inches(4.3), Inches(5.6), Inches(2.4), RGBColor(254, 243, 199), RGBColor(253, 230, 138))
    tb_r2 = slide4.shapes.add_textbox(Inches(7.133), Inches(4.5), Inches(5.2), Inches(2.0))
    tf_r2 = tb_r2.text_frame
    tf_r2.word_wrap = True
    pr3 = tf_r2.paragraphs[0]
    pr3.text = "🕒 Automated NTP Time Synchronization"
    pr3.font.size = Pt(16)
    pr3.font.bold = True
    pr3.font.color.rgb = RGBColor(180, 83, 9)
    
    pr4 = tf_r2.add_paragraph()
    pr4.text = "Eliminate clock drift forever. AutoBell automatically synchronizes hardware RTC clocks with global Network Time Protocol (NTP) servers daily."
    pr4.font.size = Pt(13)
    pr4.font.color.rgb = COLOR_DARK_TEXT
    pr4.space_before = Pt(6)

    add_footer(slide4, 4, 12)

    # ==========================================
    # SLIDE 5: PRE-ANNOUNCEMENT JINGLE & DELAY
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_bg(slide5, COLOR_SLATE_BG)
    add_header(slide5, "Audio Innovation", "Pre-Announcement Signature Chime & Pause Control")

    # Banner Card
    add_card(slide5, Inches(0.8), Inches(1.6), Inches(11.733), Inches(1.4), RGBColor(245, 243, 255), RGBColor(221, 214, 254))
    t_ban = slide5.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(11.333), Inches(1.1))
    tf_ban = t_ban.text_frame
    tf_ban.word_wrap = True
    pb1 = tf_ban.paragraphs[0]
    pb1.text = "🎵 Professional Campus Audio Identity"
    pb1.font.size = Pt(18)
    pb1.font.bold = True
    pb1.font.color.rgb = RGBColor(109, 40, 217)
    
    pb2 = tf_ban.add_paragraph()
    pb2.text = "AutoBell ensures every announcement or bell is preceded by a signature pre-announcement audio jingle, capturing student and staff attention before messages begin."
    pb2.font.size = Pt(13)
    pb2.font.color.rgb = COLOR_DARK_TEXT
    pb2.space_before = Pt(4)

    # 3 Step Cards showing Audio Flow
    steps = [
        ("STEP 1: Signature Chime", "Pre-Announcement Jingle", "Plays a high-clarity melodious chime or school anthem snippet from the global audio library."),
        ("STEP 2: Smart Pause Delay", "Custom 2–5 Second Buffer", "Configurable pause duration allows noise to settle across hallways before vocal broadcast starts."),
        ("STEP 3: Main Audio Output", "Bell / TTS / Voice Stream", "Plays scheduled bell MP3 tone, text-to-speech notice, or live administrator voice announcement.")
    ]

    for idx, (step_num, title, desc) in enumerate(steps):
        x = Inches(0.8 + idx * 4.05)
        add_card(slide5, x, Inches(3.3), Inches(3.633), Inches(3.4), COLOR_WHITE, COLOR_CARD_BORDER)
        
        # Step Top Tag
        tag = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(3.3), Inches(3.633), Inches(0.45))
        tag.fill.solid()
        tag.fill.fore_color.rgb = COLOR_BLUE_ACCENT if idx != 1 else COLOR_GOLD
        tag.line.fill.background()
        tf_t = tag.text_frame
        pt = tf_t.paragraphs[0]
        pt.text = step_num
        pt.alignment = PP_ALIGN.CENTER
        pt.font.size = Pt(11)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_WHITE

        tb = slide5.shapes.add_textbox(x + Inches(0.15), Inches(3.85), Inches(3.333), Inches(2.7))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_DARK_TEXT
        p2.space_before = Pt(10)

    add_footer(slide5, 5, 12)

    # ==========================================
    # SLIDE 6: TEXT-TO-SPEECH & LIVE BROADCASTS
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_bg(slide6, COLOR_SLATE_BG)
    add_header(slide6, "Public Address System", "Automated Text-to-Speech & Live Voice Broadcasts")

    # Left Column: TTS Feature
    add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1), COLOR_WHITE, COLOR_CARD_BORDER)
    tb_tts = slide6.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.7))
    tf_tts = tb_tts.text_frame
    tf_tts.word_wrap = True
    
    pt1 = tf_tts.paragraphs[0]
    pt1.text = "🗣️ Automated Text-to-Speech (TTS)"
    pt1.font.size = Pt(18)
    pt1.font.bold = True
    pt1.font.color.rgb = COLOR_BLUE_ACCENT

    tts_points = [
        "Instant Notice Conversion: Simply type any announcement text into the Web Dashboard or Mobile App.",
        "Natural Spoken Voices: Converts text into high-definition, natural-sounding audio announcements.",
        "Scheduled & Ad-Hoc Notices: Schedule daily morning assemblies, principal announcements, or bus alerts.",
        "Multilingual Support: Supports multi-language speech output for diverse educational environments.",
        "Zero Recording Needed: No need to pre-record voice files — type and broadcast in seconds."
    ]
    for pt in tts_points:
        p = tf_tts.add_paragraph()
        p.text = "▪ " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(8)

    # Right Column: Live Broadcast & MP3
    add_card(slide6, Inches(6.933), Inches(1.6), Inches(5.6), Inches(5.1), COLOR_WHITE, COLOR_CARD_BORDER)
    tb_bcast = slide6.shapes.add_textbox(Inches(7.133), Inches(1.8), Inches(5.2), Inches(4.7))
    tf_bc = tb_bcast.text_frame
    tf_bc.word_wrap = True
    
    pb1 = tf_bc.paragraphs[0]
    pb1.text = "🎙️ Live Voice & MP3 Audio Streaming"
    pb1.font.size = Pt(18)
    pb1.font.bold = True
    pb1.font.color.rgb = COLOR_CYAN

    bcast_points = [
        "Live Mobile Microphone Stream: Press-and-talk from administrator smartphone to campus speakers.",
        "Custom MP3 Upload Library: Upload custom school anthems, prayer Adhan calls, or event chimes.",
        "High-Fidelity Audio DAC: Supports crisp 44.1kHz audio reproduction over standard PA amplifier systems.",
        "Volume & Equalizer Control: Fine-tune speaker output levels for indoor classrooms vs outdoor grounds.",
        "Audit Trail Logging: All broadcasted announcements are logged with timestamp and user ID."
    ]
    for pt in bcast_points:
        p = tf_bc.add_paragraph()
        p.text = "▪ " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(8)

    add_footer(slide6, 6, 12)

    # ==========================================
    # SLIDE 7: HARDWARE RELIABILITY & OFFLINE-FIRST
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_bg(slide7, COLOR_SLATE_BG)
    add_header(slide7, "Industrial IoT Tech", "Offline-First Reliability & Hardware Watchdog")

    # 3 Large Feature Columns
    h_features = [
        ("💾 Offline-First Architecture", "Local LittleFS / SPIFFS Storage", "Schedules, profile definitions, and bell sound files are cached locally in ESP32 flash memory. If campus Wi-Fi or internet fails, the bell system functions seamlessly without missing a beat."),
        ("🐕 Hardware Watchdog Guard", "100% Uptime & Self-Healing", "Integrated hardware watchdog timer monitors controller CPU health. In the rare event of a power surge or freeze, the device auto-reboots and restores state in under 2 seconds."),
        ("🛡️ Zero-Downtime Guarantee", "SLA & Instant Replacement", "Supported by AutoBell's optional yearly hardware SLA. In case of hardware damage, receive an immediate pre-configured replacement unit from central inventory.")
    ]

    for idx, (title, sub, body) in enumerate(h_features):
        x = Inches(0.8 + idx * 4.05)
        add_card(slide7, x, Inches(1.6), Inches(3.633), Inches(5.1), COLOR_WHITE, COLOR_CARD_BORDER)
        
        top_bar = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.6), Inches(3.633), Inches(0.1))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = COLOR_BLUE_ACCENT if idx != 1 else COLOR_GOLD
        top_bar.line.fill.background()

        tb = slide7.shapes.add_textbox(x + Inches(0.15), Inches(1.9), Inches(3.333), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_NAVY
        
        p_sub = tf.add_paragraph()
        p_sub.text = sub
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_BLUE_ACCENT
        p_sub.space_before = Pt(4)
        
        p2 = tf.add_paragraph()
        p2.text = body
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_DARK_TEXT
        p2.space_before = Pt(14)

    add_footer(slide7, 7, 12)

    # ==========================================
    # SLIDE 8: CAMPUS EMERGENCY & LOCKDOWN SYSTEM
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_bg(slide8, COLOR_SLATE_BG)
    add_header(slide8, "Campus Safety", "Instant Emergency Alert & Lockdown Activation")

    # Big Alert Box (Red Theme)
    add_card(slide8, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.1), RGBColor(255, 241, 242), RGBColor(254, 205, 211))

    tb_em = slide8.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11.133), Inches(4.7))
    tf_em = tb_em.text_frame
    tf_em.word_wrap = True

    pe1 = tf_em.paragraphs[0]
    pe1.text = "🚨 Slide-to-Activate Campus Lockdown System"
    pe1.font.size = Pt(22)
    pe1.font.bold = True
    pe1.font.color.rgb = COLOR_RED_ACCENT

    pe2 = tf_em.add_paragraph()
    pe2.text = "In high-stakes emergencies, speed saves lives. AutoBell provides instant lockdown and disaster evacuation controls accessible directly from authorized mobile devices and web dashboards."
    pe2.font.size = Pt(14)
    pe2.font.color.rgb = COLOR_DARK_TEXT
    pe2.space_before = Pt(6)

    em_steps = [
        "🔒 Slide-to-Confirm Trigger: Prevents accidental false alarms while enabling 1-second emergency activation.",
        "🔊 Continuous Siren & Audio Override: Immediately cuts off standard bells and broadcasts high-decibel warning sirens.",
        "📱 Instant Security Notifications: Alerts principals, campus security officers, and administrators simultaneously.",
        "🛑 1-Tap Cancellation: Authorized admins can disarm lockdown and return system to normal operating state instantly.",
        "📊 Incident Time-Stamping: Records exact trigger timestamp, user identity, and duration for administrative audit compliance."
    ]

    for step in em_steps:
        p = tf_em.add_paragraph()
        p.text = "• " + step
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(10)

    add_footer(slide8, 8, 12)

    # ==========================================
    # SLIDE 9: DASHBOARD & MOBILE APP EXPERIENCE
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_bg(slide9, COLOR_SLATE_BG)
    add_header(slide9, "User Interface", "Intuitive Web Dashboard & Android Mobile App")

    # Left: Web Dashboard
    add_card(slide9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1), COLOR_WHITE, COLOR_CARD_BORDER)
    tb_web = slide9.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.7))
    tf_web = tb_web.text_frame
    tf_web.word_wrap = True
    
    pw1 = tf_web.paragraphs[0]
    pw1.text = "💻 Modern Web Dashboard"
    pw1.font.size = Pt(18)
    pw1.font.bold = True
    pw1.font.color.rgb = COLOR_BLUE_ACCENT

    web_list = [
        "URL: https://iobell.web.app/ — Accessible from any PC, tablet, or desktop browser.",
        "Drag-and-Drop Schedule Builder: Easily manage complex period timings & bell durations.",
        "Real-Time Controller Status: Live indicators for Wi-Fi signal, online status, and active profile.",
        "Multi-School Campus Management: Manage multiple school buildings or branches from one master account.",
        "Role-Based Access Control (RBAC): Super Admin, School Admin, Operator, and Partner permissions."
    ]
    for item in web_list:
        p = tf_web.add_paragraph()
        p.text = "▪ " + item
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(8)

    # Right: Mobile App
    add_card(slide9, Inches(6.933), Inches(1.6), Inches(5.6), Inches(5.1), COLOR_WHITE, COLOR_CARD_BORDER)
    tb_mob = slide9.shapes.add_textbox(Inches(7.133), Inches(1.8), Inches(5.2), Inches(4.7))
    tf_mob = tb_mob.text_frame
    tf_mob.word_wrap = True
    
    pm1 = tf_mob.paragraphs[0]
    pm1.text = "📱 Android Mobile App"
    pm1.font.size = Pt(18)
    pm1.font.bold = True
    pm1.font.color.rgb = COLOR_CYAN

    mob_list = [
        "On-The-Go Campus Control: Control bells from hallways, grounds, or off-site locations.",
        "Quick Action Widgets: 1-tap manual bell ring, voice broadcast, or emergency trigger.",
        "Push Notifications: Instant alerts for power status, schedule sync, or lockdown triggers.",
        "Lightweight & Fast: Optimized React Native Android app for instant response.",
        "Biometric Security: Quick fingerprint/face unlock for authorized school staff."
    ]
    for item in mob_list:
        p = tf_mob.add_paragraph()
        p.text = "▪ " + item
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(8)

    add_footer(slide9, 9, 12)

    # ==========================================
    # SLIDE 10: STAKEHOLDER VALUE PROPOSITION (ROI)
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_bg(slide10, COLOR_SLATE_BG)
    add_header(slide10, "Value Proposition", "Tailored Benefits for Every Campus Stakeholder")

    rows = [
        ("School Principals", "Eliminates bell chaos & manual staff dependence; 1-tap profile switching for exam/special days."),
        ("IT & Technical Directors", "Cloud dashboard management, 100% uptime with Hardware Watchdog, zero maintenance headaches."),
        ("Campus Safety Officers", "Instant mobile emergency lockdown trigger, continuous sirens, and priority public address announcements."),
        ("Faith & Specialty Schools", "Custom MP3 audio uploads for prayer calls (Adhan), custom chimes, and flexible daily schedules."),
        ("IT Solution Resellers", "High-margin SaaS recurring revenue model backed by the AutoBell Partner Portal.")
    ]

    for idx, (role, benefit) in enumerate(rows):
        y = Inches(1.6 + idx * 1.02)
        
        # Left Role Badge
        r_box = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(3.2), Inches(0.88))
        r_box.fill.solid()
        r_box.fill.fore_color.rgb = COLOR_NAVY
        r_box.line.fill.background()
        tf_r = r_box.text_frame
        pr = tf_r.paragraphs[0]
        pr.text = role
        pr.alignment = PP_ALIGN.CENTER
        pr.font.size = Pt(13)
        pr.font.bold = True
        pr.font.color.rgb = COLOR_GOLD

        # Right Benefit Card
        b_box = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.2), y, Inches(8.333), Inches(0.88))
        b_box.fill.solid()
        b_box.fill.fore_color.rgb = COLOR_WHITE
        b_box.line.color.rgb = COLOR_CARD_BORDER
        b_box.line.width = Pt(1.5)
        tf_b = b_box.text_frame
        tf_b.word_wrap = True
        pb = tf_b.paragraphs[0]
        pb.text = benefit
        pb.font.size = Pt(12)
        pb.font.color.rgb = COLOR_DARK_TEXT

    add_footer(slide10, 10, 12)

    # ==========================================
    # SLIDE 11: SECURITY, BACKUP & DATA PROTECTION
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_bg(slide11, COLOR_SLATE_BG)
    add_header(slide11, "Enterprise Architecture", "Multi-Tenant Cloud Security & Complete Data Protection")

    sec_cards = [
        ("🔒 Enterprise Data Security", "Supabase PostgreSQL & RLS", "Built on enterprise-grade Supabase cloud infrastructure with strict Row Level Security (RLS) policies ensuring isolated multi-tenant data protection."),
        ("💾 1-Click System Backup", "Full Disaster Recovery", "Complete peace of mind. Perform 1-click full backups and instant restores for database records, active schedules, diagnostic logs, and custom MP3 media files."),
        ("🔑 Role-Based Access (RBAC)", "Multi-Tiered Permissions", "Assign granular permissions to staff: Super Admin (full control), School Admin (schedules & TTS), Operator (manual ring), and Partner (reseller tracking).")
    ]

    for idx, (title, sub, desc) in enumerate(sec_cards):
        x = Inches(0.8 + idx * 4.05)
        add_card(slide11, x, Inches(1.6), Inches(3.633), Inches(5.1), COLOR_WHITE, COLOR_CARD_BORDER)
        
        top_bar = slide11.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.6), Inches(3.633), Inches(0.1))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = COLOR_BLUE_ACCENT if idx != 1 else COLOR_CYAN
        top_bar.line.fill.background()

        tb = slide11.shapes.add_textbox(x + Inches(0.15), Inches(1.9), Inches(3.333), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_NAVY
        
        p_sub = tf.add_paragraph()
        p_sub.text = sub
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_BLUE_ACCENT
        p_sub.space_before = Pt(4)
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_DARK_TEXT
        p2.space_before = Pt(14)

    add_footer(slide11, 11, 12)

    # ==========================================
    # SLIDE 12: CALL TO ACTION & CLOSING (Dark Theme)
    # ==========================================
    slide12 = prs.slides.add_slide(blank_layout)
    set_bg(slide12, COLOR_NAVY)

    # Center Container Card
    add_card(slide12, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.0), RGBColor(30, 41, 59), COLOR_BLUE_ACCENT)

    tb_c = slide12.shapes.add_textbox(Inches(1.8), Inches(1.5), Inches(9.733), Inches(4.4))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    pc1 = tf_c.paragraphs[0]
    pc1.text = "Upgrade Your Campus Operations Today"
    pc1.alignment = PP_ALIGN.CENTER
    pc1.font.size = Pt(32)
    pc1.font.bold = True
    pc1.font.color.rgb = COLOR_WHITE

    pc2 = tf_c.add_paragraph()
    pc2.text = "Transform your school bell schedules, public announcements, and emergency safety with AutoBell."
    pc2.alignment = PP_ALIGN.CENTER
    pc2.font.size = Pt(16)
    pc2.font.color.rgb = RGBColor(203, 213, 225)
    pc2.space_before = Pt(10)

    # CTA Buttons / Highlights
    pc3 = tf_c.add_paragraph()
    pc3.text = "\n🚀 Experience the Live Dashboard:"
    pc3.alignment = PP_ALIGN.CENTER
    pc3.font.size = Pt(18)
    pc3.font.bold = True
    pc3.font.color.rgb = COLOR_GOLD

    pc4 = tf_c.add_paragraph()
    pc4.text = "https://iobell.web.app/"
    pc4.alignment = PP_ALIGN.CENTER
    pc4.font.size = Pt(22)
    pc4.font.bold = True
    pc4.font.color.rgb = COLOR_CYAN

    pc5 = tf_c.add_paragraph()
    pc5.text = "\n📞 Contact Us for Live Demo & On-Site Installation Options"
    pc5.alignment = PP_ALIGN.CENTER
    pc5.font.size = Pt(14)
    pc5.font.color.rgb = COLOR_WHITE

    add_footer(slide12, 12, 12, dark_mode=True)

    # Save output
    output_path = r"c:\My Drive\My Drive\AutoBell\AutoBell_Customer_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_presentation()
